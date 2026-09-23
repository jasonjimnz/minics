"""Convert imported documents to clean markdown.

Every file the user imports becomes markdown.  That markdown is what the RAG
chunks, what the graph extracts from, and what the editor shows.  Keeping one
canonical text form means grounding always points at something the user can read
and edit.

Supported inputs: PDF, plain text, Markdown, DOCX and LaTeX.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS: dict[str, str] = {
    ".pdf": "pdf",
    ".txt": "text",
    ".text": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".docx": "docx",
    ".tex": "latex",
    ".latex": "latex",
    ".ltx": "latex",
}

_ZERO_WIDTH = re.compile(r"[\u200b\u200c\u200d\ufeff]")
_MULTI_BLANK = re.compile(r"\n{3,}")
_TRAILING_WS = re.compile(r"[ \t]+\n")
_PAGE_NUMBER = re.compile(r"^\s*(?:page\s*)?\d{1,4}\s*(?:/\s*\d{1,4})?\s*$", re.IGNORECASE)
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")


class UnsupportedDocumentError(ValueError):
    """Raised for a file type MiniCS cannot convert."""


@dataclass
class ConvertedDocument:
    kind: str
    title: str
    markdown: str
    metadata: dict


def detect_kind(filename: str | Path) -> str:
    suffix = Path(filename).suffix.lower()
    kind = SUPPORTED_EXTENSIONS.get(suffix)
    if kind is None:
        raise UnsupportedDocumentError(
            f"Unsupported file type '{suffix or '(none)'}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return kind


def is_supported(filename: str | Path) -> bool:
    return Path(filename).suffix.lower() in SUPPORTED_EXTENSIONS


def read_text(path: str | Path) -> str:
    data = Path(path).read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")  # pragma: no cover


# ---------------------------------------------------------------------------
# Markdown cleanup
# ---------------------------------------------------------------------------
def clean_markdown(markdown: str) -> str:
    """Deterministic, dependency-free tidy-up of converted markdown."""
    if not markdown:
        return ""
    text = markdown.replace("\r\n", "\n").replace("\r", "\n")
    text = _ZERO_WIDTH.sub("", text)
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    lines = []
    for line in text.split("\n"):
        stripped = line.rstrip()
        if _PAGE_NUMBER.match(stripped):
            continue
        lines.append(stripped)
    text = "\n".join(lines)
    text = _TRAILING_WS.sub("\n", text)
    text = _MULTI_BLANK.sub("\n\n", text)
    return text.strip() + "\n" if text.strip() else ""


# ---------------------------------------------------------------------------
# Per-format converters
# ---------------------------------------------------------------------------
def pdf_to_markdown(path: str | Path) -> tuple[str, dict]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    meta = {}
    try:
        info = reader.metadata or {}
        meta = {
            "author": getattr(info, "author", None) or info.get("/Author", ""),
            "title": getattr(info, "title", None) or info.get("/Title", ""),
        }
    except Exception:  # noqa: BLE001 - metadata is best-effort
        meta = {}
    pages = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:  # noqa: BLE001 - keep going on a bad page
            text = ""
        pages.append(text.strip())
    markdown = "\n\n".join(p for p in pages if p)
    meta["pages"] = len(reader.pages)
    return markdown, meta


def docx_to_markdown(path: str | Path) -> tuple[str, dict]:
    try:
        import mammoth
        from markdownify import markdownify

        with open(path, "rb") as handle:
            result = mammoth.convert_to_html(handle)
        raw = markdownify(
            result.value,
            heading_style="ATX",
            bullets="-",
            strip=["img"],
        )
        return raw, {"messages": [m.message for m in getattr(result, "messages", [])][:10]}
    except Exception:  # noqa: BLE001 - fall back to python-docx
        from docx import Document as DocxDocument

        document = DocxDocument(str(path))
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "").lower()
            if style.startswith("heading"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                blocks.append(f"{'#' * min(int(level), 6)} {text}")
            elif style.startswith("list"):
                blocks.append(f"- {text}")
            else:
                blocks.append(text)
        for table in document.tables:
            rows = [
                "| " + " | ".join(cell.text.strip() for cell in row.cells) + " |"
                for row in table.rows
            ]
            if rows:
                header = rows[0]
                separator = "| " + " | ".join("---" for _ in table.rows[0].cells) + " |"
                blocks.append("\n".join([header, separator, *rows[1:]]))
        return "\n\n".join(blocks), {"fallback": "python-docx"}


def text_to_markdown(path: str | Path) -> tuple[str, dict]:
    text = read_text(path)
    lines = text.replace("\r\n", "\n").split("\n")
    converted: list[str] = []
    for line in lines:
        stripped = line.strip()
        if stripped and len(stripped) < 80 and stripped.isupper() and stripped.isascii():
            converted.append(f"## {stripped.title()}")
        else:
            converted.append(line)
    return "\n".join(converted), {}


def markdown_to_markdown(path: str | Path) -> tuple[str, dict]:
    return read_text(path), {}


_TEX_REPLACEMENTS: tuple[tuple[str, str], ...] = (
    (r"\\documentclass(\[[^\]]*\])?\{[^}]*\}", ""),
    (r"\\usepackage(\[[^\]]*\])?\{[^}]*\}", ""),
    (r"\\begin\{document\}", "\n\n"),
    (r"\\end\{document\}", "\n\n"),
    (r"\\title\{([^}]*)\}", r"\n\n# \1\n\n"),
    (r"\\part\{([^}]*)\}", r"\n\n# \1\n\n"),
    (r"\\chapter\{([^}]*)\}", r"\n\n# \1\n\n"),
    (r"\\section\{([^}]*)\}", r"\n\n## \1\n\n"),
    (r"\\subsection\{([^}]*)\}", r"\n\n### \1\n\n"),
    (r"\\subsubsection\{([^}]*)\}", r"\n\n#### \1\n\n"),
    (r"\\paragraph\{([^}]*)\}", r"\n\n##### \1\n\n"),
    (r"\\textbf\{([^}]*)\}", r"**\1**"),
    (r"\\textit\{([^}]*)\}", r"*\1*"),
    (r"\\emph\{([^}]*)\}", r"*\1*"),
    (r"\\texttt\{([^}]*)\}", r"`\1`"),
    (r"\\underline\{([^}]*)\}", r"\1"),
    (r"\\begin\{itemize\}", "\n"),
    (r"\\end\{itemize\}", "\n"),
    (r"\\begin\{enumerate\}", "\n"),
    (r"\\end\{enumerate\}", "\n"),
    (r"\\item\b", "\n- "),
    (r"\\cite\{([^}]*)\}", r"[\1]"),
    (r"\\ref\{([^}]*)\}", r"(\1)"),
    (r"\\label\{([^}]*)\}", ""),
    (r"\\footnote\{([^}]*)\}", r" (\1)"),
    (r"\\%", "%"),
    (r"\\&", "&"),
    (r"\\_", "_"),
    (r"~", " "),
)


def latex_to_markdown(text: str) -> tuple[str, dict]:
    # Protect math so the text conversion pass does not mangle it.
    math_store: list[str] = []

    def _stash(match: re.Match[str]) -> str:
        math_store.append(match.group(0))
        return f"@@MATH{len(math_store) - 1}@@"

    protected = re.sub(
        r"(\$\$.*?\$\$|\\\[.*?\\\]|\\\(.*?\\\)|\$[^$\n]+?\$)", _stash, text, flags=re.S
    )
    result = protected
    for pattern, replacement in _TEX_REPLACEMENTS:
        result = re.sub(pattern, replacement, result)
    try:  # strip any remaining commands with pylatexenc
        from pylatexenc.latex2text import LatexNodes2Text

        result = LatexNodes2Text().latex_to_text(result)
    except Exception:  # noqa: BLE001 - regex pass already did the heavy lifting
        result = re.sub(r"\\[a-zA-Z@]+\*?(\[[^\]]*\])?(\{[^}]*\})?", "", result)
    for index, value in enumerate(math_store):
        result = result.replace(f"@@MATH{index}@@", value)
    return result, {}


def convert_file(path: str | Path) -> ConvertedDocument:
    """Convert ``path`` to markdown, returning the text plus light metadata."""
    path = Path(path)
    kind = detect_kind(path)
    if kind == "pdf":
        markdown, meta = pdf_to_markdown(path)
    elif kind == "docx":
        markdown, meta = docx_to_markdown(path)
    elif kind == "latex":
        markdown, meta = latex_to_markdown(read_text(path))
    elif kind == "markdown":
        markdown, meta = markdown_to_markdown(path)
    else:
        markdown, meta = text_to_markdown(path)

    title = _guess_title(markdown, meta) or path.stem.replace("_", " ").replace("-", " ").strip()
    meta["kind"] = kind
    return ConvertedDocument(
        kind=kind,
        title=title,
        markdown=clean_markdown(markdown),
        metadata=meta,
    )


def _guess_title(markdown: str, meta: dict) -> str:
    for key in ("title",):
        value = (meta or {}).get(key)
        if value:
            return str(value).strip()
    for line in (markdown or "").splitlines():
        stripped = line.strip()
        if stripped.startswith("# "):
            return stripped[2:].strip()
        if stripped:
            return ""
    return ""
